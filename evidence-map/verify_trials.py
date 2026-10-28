from pathlib import Path
import csv
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
p=Path('user/workspace/co-scientist-outputs/implantable-device-evidence')
w=load_workbook(p/'implantable_device_evidence.xlsx')
urls={
 'Q':'https://pmc.ncbi.nlm.nih.gov/articles/PMC5126155/',
 'AVP':'https://doi.org/10.1056/NEJMoa2300080',
 'BSP':'https://doi.org/10.1056/NEJMoa2401807',
 'BIP':'https://pmc.ncbi.nlm.nih.gov/articles/PMC10844377/',
 'LON':'https://pmc.ncbi.nlm.nih.gov/articles/PMC12330253/',
 'ALG':'https://pmc.ncbi.nlm.nih.gov/articles/PMC12367825/',
 'MAR':'https://portal.findresearcher.sdu.dk/en/publications/atrioventricular-synchronous-pacing-using-a-leadless-ventricular-/',
 'ACC':'https://ohsu.elsevierpure.com/en/publications/ambulatory-atrioventricular-synchronous-pacing-over-time-using-a-/',
 'IDE':'https://ohiostate.elsevierpure.com/en/publications/a-leadless-intracardiac-transcatheter-pacing-system/',
 'RAV':'https://clinicaltrials.gov/study/NCT05252702',
 'RBS':'https://clinicaltrials.gov/study/NCT04798768'}
# One record per primary report; exact device matching is deliberately separate from study quality.
th=['Study ID','Candidate ID','Primary report','Publication date','Registry ID','Design / source material','Actual model tested','Relation to candidate','Population / follow-up','Source access','Key verification','GRADE-lite scope','Funding / independence','Timing versus review','URL']
trials=[
 ['Q','QIN','Xiang et al.; Qinming8631 DR trial; DOI 10.4103/0366-6999.193443','2016-11-20','Not identified in checked paper','Randomized, multicenter, single-blind noninferiority trial','Qinming8631 DR versus Talos DR','Direct model match','225 randomized (113/112); 14 centers; 6 months','Full HTML methods, tables and results','FAS 92/113 vs 89/112; PPS 83/87 vs 77/86. QM7211 is ventricular lead.','Low for 6-month effective-pacing endpoint; not AI benefit','Paper reports no funding and no conflicts','Before review',urls['Q']],
 ['AVP','AVE','Knops et al.; A Dual-Chamber Leadless Pacemaker; DOI 10.1056/NEJMoa2300080','2023-05-20','NCT05252702','Prospective international single-arm pivotal trial','AVEIR DR','Direct model match','First 300 attempted implants; 55 centers; 90-day safety / 3-month performance','Full original NEJM PDF; Table 3 and Figure 2 visually checked','Implant 295/300; safety 271/300; performance endpoints use imputation and differing analysis populations','Low for short-term descriptive endpoints, including quantified safety','Abbott funded; sponsor analyses; independent event adjudication','Before review',urls['AVP']],
 ['BSP','BOS','Knops et al.; A Modular Communicative Leadless Pacing–Defibrillator System; DOI 10.1056/NEJMoa2401807','2024-05-18 (updated 2024-05-22)','NCT04798768','Prospective international single-group; prespecified early analysis','EMBLEM S-ICD + EMPOWER','Direct system match; later publication','293 enrolled; 162 endpoint cohort; 151 completed 6-month visit; median follow-up 12.4 months','Full updated NEJM PDF; Figures 2–3 visually checked','97.5% LP-major-complication-free; 98.8% modeled communication; 147/151 capture responders; ATP 61.3% of 31 episodes','Low for primary descriptive endpoints; Very low for sparse ATP termination estimate','Boston Scientific funded; sponsor analyses; independent event adjudication','After February 2024 review',urls['BSP']],
 ['BIP','BIO','Bisignani et al.; SmartECG remote ECG review validation; DOI 10.3389/fcvm.2024.1343424','2024-01-23','Pooled CERTITUDE, NCT04025710, NCT04075084','Retrospective replay of recordings pooled from clinical projects; separate time-motion sample','BioMonitor IIIm recordings processed by SmartECG for BioMonitor IV workflow','Proxy hardware, balanced setting; not the 86%/98% Very Specific dataset','368 patients; 418.5 patient-years; 61,517 sECGs; 2,000 adjudicated sample','Full primary HTML including algorithm, methods, tables and funding','42.8% recordings filtered; 15/708 true episodes rejected; weighted rejection 2.56%; algorithm resides on remote platform','Low for replay results; Very low when transferred to prospective BioMonitor IV clinical benefit','Biotronik funded; manufacturer employees coauthor; NOT independent external replication','Before cover date; after review reference retrieval date 2024-01-16',urls['BIP']],
 ['MAR','MAV','Steinwender et al.; MARVEL 2; DOI 10.1016/j.jacep.2019.10.017','2020-01 (journal issue)','NCT03752151','Prospective acute within-patient VDD versus VVI-50 evaluation','Micra with downloaded investigational algorithm, not AV2','Earlier generation / algorithm feasibility','75 enrolled at 12 centers; efficacy subgroup 40; up to 5 hours','Original article abstract reproduced in author institution repository; full text not audited','95% vs 0% achieved ≥70% resting AV synchrony; mean 89.2% vs 26.8%','Low for acute earlier-generation algorithm performance; Very low for AV2 transfer','Manufacturer technical involvement; full funding/declarations not audited','Before review; not a direct AV2 trial',urls['MAR']],
 ['ACC','MAV','Chinitz et al.; AccelAV; DOI 10.1016/j.hrthm.2022.08.033','2023-01 (journal issue)','NCT04245345','Prospective single-arm clinical study','Micra AV (first generation)','Earlier-generation clinical context; not AV2','152 implanted; 20 centers; primary subgroup 54; optimize subgroup 20; through 3 months','Original abstract plus funding/declarations reproduced in author institution repository; full text not audited','Mean resting AVS 85.4%; ambulatory 74.8%; optimized subgroup 82.6%','Low for select Micra AV short-term endpoints; Very low for AV2 transfer','Medtronic funded; manufacturer employee/shareholder coauthors','Before review; not a direct AV2 trial',urls['ACC']],
 ['IDE','MVR','Reynolds et al.; Micra transcatheter pacing study; DOI 10.1056/NEJMoa1511643','2016-02-11 (journal issue)','NCT02004873','Prospective single-arm; historical performance goals; post-hoc historical control','Original Micra ventricular device, not VR2','Earlier generation; not direct VR2 outcomes','725 attempted implants; efficacy subset 297 paired 6-month records','Original abstract reproduced in author institution repository; full text not audited','719/725 implant success; safety KM 96.0%; efficacy 292/297 (98.3%)','Low for first-generation descriptive endpoints; Very low for VR2 transfer','Full funding/declarations not audited; manufacturer-related study context','Before review; not a direct VR2 trial',urls['IDE']],
 ['LON','MAV; MVR','Leal et al.; Device longevity of a leadless pacemaker family; DOI 10.1080/14796678.2025.2527466','2025-07-16','No new device-randomized trial','Longevity and patient-survival modeling; Monte Carlo simulation','AV2/VR2 battery/electronics model using AV/VR pacing data','Direct design model, not observed long-term AV2/VR2 outcomes','644 VR IDE input patients; 999 AV CareLink input patients; historical transvenous survival','Full original HTML, methods, limitations and funding','AV2 15.6 years; VR2 16.7 years projected; lifetime single device 80% and 91%','Very low for actual lifetime clinical outcomes (modeled extrapolation)','Medtronic funded; employee/shareholder coauthors; underlying data unavailable','After review; cannot retroactively support its 2024 evidence base',urls['LON']],
 ['ALG','MAV','Leal et al.; AV synchrony virtual patient analysis; DOI 10.1007/s13239-025-00783-0','2025-05-27','Uses AccelAV data; no new comparative AV2 clinical trial','Bench/Monte Carlo simulation plus anonymous field-representative survey','AV versus AV2 algorithms driven by virtual signals from Micra AV patients','AV2 algorithm modeling, not prospective patient efficacy trial','30 virtual setup cases; 20 virtual ambulatory cases; surveys AV 104 / AV2 98','Full original HTML methods, results, limitations and funding','Simulated mean AVS 91.9% vs 68.0% after setup; 84.1% vs 49.5% at 80–100 beats/min','Very low for clinical benefit; survey estimate not measured time-motion outcome','Medtronic funded; employee/shareholder coauthors; reused model-calibration data','After review',urls['ALG']]
]
# Add endpoint-level primary verification while retaining all original claims unchanged.
ws=w['Evidence'];oldheads=[c.value for c in ws[1]]
extra=['Study ID','Actual model / evidence type','Publication date','Timing versus review','Uncertainty / CI','Verification status']
for j,h in enumerate(extra,len(oldheads)+1):ws.cell(1,j,h)
for i in range(2,ws.max_row+1):
 cid=ws.cell(i,2).value;ws.cell(i,ws.max_column,'Original extraction retained; see primary verification rows')
new=[]
def ev(sid,cid,endpoint,text,val=None,unit='%',n='NR',follow='NR',comp='No concurrent control',setting='NR',grade='Low',reason='Single-arm descriptive performance; no randomized comparative or AI effect inference',loc='',note='',ci='Not reported for this estimate',model='',status='Verified in primary full text'):
 t=next(x for x in trials if x[0]==sid)
 base=[f'V{len(new)+1:02}',cid,endpoint,text,'=' if val is not None else '',val,None,unit,'Primary-study verification',n,follow,comp,setting,grade,reason,loc,sid,urls[sid],note]
 new.append(base+[sid,model or t[6],t[3],t[13],ci,status])
ev('Q','QIN','Effective pacing, FAS','92/113 (81.4%) vs 89/112 (79.5%)',81.4,n='113 Qinming /112 Talos',follow='6 months',comp='Talos DR, randomized',setting='Noninferiority margin −9.5 percentage points; LOCF then missing=ineffective',reason='RCT High −1 missing endpoints/protocol deviations −1 imprecision/reporting limitations = Low',loc='Methods; Table 3; Results',note='Confirm group rates; reported adjusted differences cannot be treated as crude differences. Not proof of superiority.')
ev('Q','QIN','Effective pacing, PPS','83/87 (95.4%) vs 77/86 (89.5%)',95.4,n='87 Qinming /86 Talos',follow='6 months',comp='Talos DR',reason='Post-randomization exclusions plus inconsistent reported differences limit inference',loc='Tables 2–3',note='52 of 225 randomized not in PPS; published PPS difference is 1.8 in table versus 11.8 in prose; do not pool with FAS.')
ev('AVP','AVE','Implantation success','295/300 (98.3%)',98.3,n='300 attempted implants',follow='Implant procedure',loc='PDF pp.1,5',note='Requires two functioning implanted devices AND established communication; two atrial nonimplants and three inadequate communication cases.')
ev('AVP','AVE','Freedom from device/procedure-related serious adverse events','271/300 (90.3%)',90.3,n='300 attempted implants',follow='90 days',comp='Performance goal 78%',loc='PDF pp.4–5; Table 3 p.7',ci='Two-sided 95% CI 87.0–93.7%',note='35 adjudicated complications in 29 patients; safety goal achievement is not absence of serious events.')
ev('AVP','AVE','Adequate atrial capture AND sensing','90.2% after imputation',90.2,n='299 analysis; 290 measurable; 2 implant failures; 7 missing imputed',follow='3 months',comp='Performance goal 82.5%',setting='Capture ≤3.0 V at 0.4 ms AND P-wave ≥1.0 mV',loc='PDF pp.4,7,9',ci='Two-sided 95% CI 86.8–93.6%',note='Do not replace with registry 90.8% (different analysis population/exclusions).')
ev('AVP','AVE','Patients meeting seated AV-synchrony criterion','97.3% after imputation',97.3,n='294 analysis; 277 measurable; 2 implant failures; 15 missing imputed',follow='3 months',comp='Performance goal 83%',setting='≥70% evaluable cycles; ventricular event within 300 ms of atrial event; 5-minute seated recording',loc='PDF pp.4,9',ci='Two-sided 95% CI 95.4–99.3%',note='NOT 97.3% mean synchronous beats and NOT all-day ambulatory AVS; complete-case registry value 98.2% is separate.')
for posture,val,n,lo,hi in [('Sitting',98.0,277,97.0,99.1),('Supine',96.5,275,95.3,97.8),('Left lateral recumbent',97.1,275,95.7,98.6),('Right lateral recumbent',97.3,273,96.1,98.5),('Standing',98.3,274,97.6,99.1),('Normal walk',98.4,234,97.6,99.2),('Fast walk',97.7,190,96.6,98.9)]:
 ev('AVP','AVE','Mean AV synchrony: '+posture,f'{val}%',val,n=str(n)+' measurable patients',follow='3-month visit',setting='Supervised posture/gait recording',loc='PDF Figure 2A p.8, visually checked',ci=f'95% CI {lo}–{hi}%; not adjusted for multiplicity',note='Condition-specific average; not ≥70%-success proportion or continuous real-world follow-up.')
ev('BSP','BOS','Freedom from major leadless-pacemaker-related complications','97.5%',97.5,n='162 prespecified early-analysis cohort',follow='6 months',comp='Performance goal 86%',loc='PDF pp.3–6; Figure 2 p.8',ci='One-sided 98.8% lower bound 94.2% with competing risks; 92.5% without competing risks',note='Four major LP-related complications. Not freedom from ALL system complications; registry shows 92.5% bound using other analysis.')
ev('BSP','BOS','Overall system-related complication-free estimate','91.4%',91.4,n='162 endpoint cohort',follow='6 months',loc='PDF Discussion p.8',note='Broader safety definition than 97.5% LP-specific primary endpoint; do not rank against AVEIR safety definitions.')
ev('BSP','BOS','Successful wireless communication','98.8% modeled test success',98.8,n='147 patients with testing per paper; posture repeated measures',follow='6-month visit',comp='Performance goal 88%',setting='Four body postures; repeated-measures logistic regression',loc='PDF pp.4,6; Figure 3A–B p.9',ci='One-sided 97.5% lower bound 97.0%',note='Registry later lists 149 participants/589 postures; figure shows posture-specific n=145/145/146/145. Preserve original-publication population.')
ev('BSP','BOS','Capture threshold responders','147/151 (97.4%)',97.4,n='151 measured patients',follow='6-month visit',comp='Performance goal 80%',setting='Capture threshold ≤2.0 V at 0.4 ms',loc='PDF p.6; Figure 3C p.9',ci='One-sided 97.5% lower bound 93.4%')
ev('BSP','BOS','ATP termination of spontaneous arrhythmia episodes','61.3% of 31 treated episodes',61.3,n='31 episodes in 13 patients',follow='Available follow-up; median cohort follow-up 12.4 months',setting='ATP after S-ICD request; sensed arrhythmia rate 158–261 beats/min',grade='Very low',reason='Sparse episodes clustered within 13 selected patients; exploratory, no control',loc='PDF pp.3,6–7',note='Three episodes accelerated by ATP; not 61.3% of all enrolled patients. Not restricted to a universal fixed 6-month observation window.')
ev('BIP','BIO','sECGs filtered by SmartECG','26,324/61,517 (42.8%)',42.8,n='368 patients; 61,517 sECGs; 418.5 patient-years',follow='Pooled recordings through 2022-10-14',comp='Same recordings without SmartECG filtering',setting='Balanced AF setting; HMSC remote software; AI AF only',reason='Low for retrospective replay; prospective BioMonitor IV clinical transfer Very low',loc='Methods; Table 2; Results',note='Not the 86% false-positive reduction claim; denominator is ALL transmitted sECGs; IIIm proxy recordings, not prospectively implanted IV cohort.')
ev('BIP','BIO','True episodes rejected, adjudicated sample','15/708 (2.1%)',2.1,n='708 adjudicated true events among 2,000 sampled sECGs',follow='Retrospective adjudication',setting='500 samples each AF/brady/pause/tachy; balanced setting',reason='Low for sampled recording performance; patient-important diagnostic harm not evaluated',loc='Sensitivity; Table 4',note='Unweighted sample fraction, not 2.56% distribution-weighted estimate; abstract rounds weighted result to 2.6%. Narrative incorrectly says zero bradycardia rejects; Table 4 has six.')
ev('BIP','BIO','Distribution-weighted true episodes rejected','2.56%',2.56,n='Weighted from 2,000 adjudicated sECGs',follow='Retrospective replay',setting='Weight by full data-set arrhythmia distribution; balanced setting',reason='Low for recorded-signal estimate; not prospective clinical sensitivity',loc='Table 4; Results',note='Not direct reproduction of 98% preservation under Very Specific filter.')
ev('BIP','BIO','Estimated annual review workload','83 min without vs 47.7 min with SmartECG',47.7,unit='min/patient-year',n='9 reviewers, 5 hospitals; recording frequency from 368 patients',follow='Modeled annual workload',setting='Assumes reviewers ignore filtered sECGs; excludes patient-contact time',grade='Very low',reason='Model based on limited time-motion observations and unobserved alert settings',loc='Workload estimation; Results',note='Not a randomized real-world workflow/time saving outcome.')
ev('MAR','MAV','Patients with ≥70% resting AV synchrony','95% VDD vs 0% VVI-50',95,n='40 of 75 enrolled; complete AV block and sinus rhythm',follow='Acute, up to 5 hours',comp='Within-patient VVI-50',loc='Primary article abstract',grade='Very low',reason='Earlier algorithm/device and acute selected subgroup; indirect for AV2 clinical benefit',note='Primary trial verified, but this is NOT Micra AV2. MARVEL 2 numeral does not denote AV2 generation.',status='Verified in original abstract via author institution')
ev('MAR','MAV','Mean resting AV synchrony','89.2% VDD vs 26.8% VVI',89.2,n='40-patient efficacy subgroup',follow='Acute, up to 5 hours',comp='Within-patient VVI',loc='Primary article abstract',grade='Very low',reason='Indirect generation and acute endpoint for AV2',note='Mean synchronous beats differs from percentage of patients meeting ≥70% criterion.',status='Verified in original abstract via author institution')
for metric,val,n in [('Resting AV synchrony',85.4,'54 with normal sinus function and complete AV block'),('Ambulatory AV synchrony',74.8,'54 with normal sinus function and complete AV block'),('Optimized ambulatory AV synchrony',82.6,'20 programming-optimization subgroup')]:
 ev('ACC','MAV',metric,f'{val}%',val,n=n,follow='Primary assessment 1 month; study through 3 months',setting='First-generation Micra AV; programming context differs by subgroup',grade='Very low',reason='Low for Micra AV source endpoint; additional generation indirectness for AV2',loc='Primary article abstract',note='Do not attribute to AV2 or pool subgroup with all implants.',status='Verified in original abstract via author institution')
ev('IDE','MVR','Implant success, original Micra','719/725 (99.2%)',99.2,n='725 attempts',follow='Implant',grade='Very low',reason='First-generation results are indirect for VR2',loc='Primary article abstract',note='Not a VR2 implant-success trial.',status='Verified in original abstract via author institution')
ev('IDE','MVR','Freedom from major complications, original Micra','96.0% Kaplan–Meier estimate',96,n='725 enrolled; time-to-event estimate',follow='6 months',comp='Performance goal 83%',grade='Very low',reason='Original device, historical goals, no randomized control; indirect for VR2',loc='Primary article abstract',ci='95% CI 93.9–97.3%',note='Do not substitute crude 700/725; Kaplan–Meier estimand.',status='Verified in original abstract via author institution')
for cid,val,n in [('MAV',15.6,'999 first-generation Micra AV CareLink inputs'),('MVR',16.7,'644 first-generation Micra VR IDE inputs')]:
 ev('LON',cid,'Median projected longevity to recommended replacement time',f'{val} years',val,unit='years',n=n,follow='Simulated lifetime, not observed follow-up',comp='AV or VR with same underlying patient pacing inputs',setting='Stochastic battery capacity/current drain plus pacing parameters; historical patient-survival model',grade='Very low',reason='Model-based extrapolation, not observed AV2/VR2 lifetime clinical performance',loc='Methods §2; Results §3.1; Limitations',note='Later primary modeling explains approximate 16/17-year claims; cannot retroactively validate 2024 evidence. No new 16-year clinical trial.')
ev('ALG','MAV','Mean AV synchrony after automatic setup, virtual patients','91.9% AV2 vs 68.0% AV',91.9,n='30 virtual patient cases derived from AccelAV',follow='30-minute simulated resting period',comp='Simulated Micra AV algorithm',setting='No manual programming; signal-distribution calibrated model',grade='Very low',reason='Bench/Monte Carlo outcome using reused calibration data; indirect for patient clinical efficacy',loc='Methods; Results',ci='AV2 95% CI 83.5–96.2%; AV 60.7–74.6%',note='Not a randomized implanted-patient trial.')
ev('ALG','MAV','Mean AV synchrony at high sinus rates, virtual patients','84.1% AV2 vs 49.5% AV',84.1,n='20 virtual ambulatory cases',follow='Simulation from up to 24-hour source recordings',comp='Simulated Micra AV Auto A3 Threshold',setting='80–100 beats/min; upper tracking rate 105 beats/min; PVAB 550 ms; A3 window end 700–800 ms',grade='Very low',reason='Simulation, simplified physiology and reused calibration data',loc='Automated A3 Threshold Algorithm Simulation; Results',ci='AV2 95% CI 71.3–91.8%; AV 43.2–55.8%',note='Different from hardware tracking capability up to 135 beats/min.')
for r in new:ws.append(r)
# Registry values retained separately, never substituted into original-publication denominators.
rh=['Check ID','Study','Published primary report','Current registry record','Interpretation / action','Registry last update','Registry URL']
recon=[
 ['R01','AVEIR DR','300 attempted implants; 90.3% safety; 90.2% atrial performance (n=299 analysis); 97.3% synchrony (n=294 analysis)','464 enrolled overall; original first-300 safety 90.3%; atrial 90.8% (n=297); synchrony 98.2% (n=277)','Different enrollment cutoffs, exclusions and missing-data estimands. Keep publication imputed estimates distinct from registry evaluable/exclusion-based estimates; no silent replacement.','2026-05-27',urls['RAV']],
 ['R02','MODULAR ATP enrollment','293 enrolled at Jan 24 2024 cutoff; 162 endpoint cohort','297 eventual enrolled; 286 implanted; 162 interim cohort','Different cutoff/cohort, not proof that 293 is erroneous. Do not use 297 as endpoint denominator.','2026-10-15',urls['RBS']],
 ['R03','MODULAR ATP communication','147 patients with 6-month tests; figure posture denominators 145/145/146/145; modeled 98.8%','149 participants, 589 posture tests; modeled 98.8%','Participant/test totals differ; exact patient-level reason unresolved without sponsor data. Flag rather than merge; primary numerical estimate agrees.','2026-10-15',urls['RBS']],
 ['R04','MODULAR ATP safety confidence bound','97.5% with one-sided 98.8% lower bound 94.2% (competing risk); sensitivity bound 92.5% without competing risk','97.5%; one-sided 98.8% lower bound 92.5%','Both bounds appear in publication under different estimators. Name estimator and confidence level; 94.2 is not a two-sided 95% lower limit.','2026-10-15',urls['RBS']]]

def add_sheet(name,heads,rows):
 s=w.create_sheet(name);s.append(heads)
 for r in rows:s.append(r)
 s.add_table(Table(displayName=name.replace(' ','')+'Table',ref=s.dimensions))
 return s
add_sheet('Primary trials',th,trials);add_sheet('Registry checks',rh,recon)
ss=w['Sources']
for t in trials:
 if t[0]!='Q':ss.append([t[0],t[2],t[14],t[9],t[10]+' | '+t[13]])
ss.append(['RAV','ClinicalTrials.gov NCT05252702',urls['RAV'],'API snapshot checked 2026-10-28; record updated 2026-05-27','Current record; primary n=300 separate from eventual enrollment n=464'])
ss.append(['RBS','ClinicalTrials.gov NCT04798768',urls['RBS'],'API snapshot checked 2026-10-28; record updated 2026-10-15','Current record; primary published and registry population differences flagged'])
cs=w['Candidates']; cs.cell(1,15,'Primary verification / evidence timing')
cs.cell(1,6,'Original review metrics / claims');cs.cell(1,10,'Original extraction settings');cs.cell(1,11,'Original extraction limitations')
cs.cell(1,16,'Primary source IDs');cs.cell(1,17,'Primary source URL')
updates={
 'BIO':('Low (replay only)','42.8% all-recording reduction in IIIm replay at Balanced setting; 86%/98% remains unverified data-on-file claim.','Low for recorded-signal validation; prospective IV clinical benefit Very low. AI in remote platform, not implanted processor.'),
 'MAV':('Very low','No direct AV2 prospective efficacy trial identified in targeted check. MARVEL 2/AccelAV are earlier generation; 2025 papers are simulations.','15.6-year modeled longevity and virtual AVS are later evidence, not observed AV2 long-term patient outcomes.'),
 'MVR':('Very low','Original Micra IDE is not VR2. Later longevity model projects 16.7 years, not observed VR2 durability.','Do not transfer 99.2% original-Micra implant success or 96.0% KM safety directly to VR2.'),
 'QIN':('Low','Randomized six-month trial verified; FAS and PPS group rates confirmed; model identity corrected.','Reporting inconsistencies remain; no AI trial.'),
 'BOS':('Low (later evidence)','MODULAR ATP primary publication available May 2024, after review; n=162 early endpoint cohort.','97.5% LP-specific safety; broader system safety 91.4%; 98.8% communication; 97.4% capture responders.'),
 'AVE':('Low','Primary 2023 trial verified: implant 295/300; safety 271/300; explicit AVS criterion; denominators/imputation recovered.','Safety quantified at 90.3% through 90 days; primary AVS success 97.3% is ≥70% cycles during five-minute seated test.')}
for i in range(2,cs.max_row+1):
 cid=cs.cell(i,1).value;grade,check,note=updates[cid]
 cs.cell(i,7,grade);cs.cell(i,15,check+' '+note)
 primary={'BIO':('BIP','BIP'),'MAV':('MAR; ACC; ALG; LON','ALG'),'MVR':('IDE; LON','LON'),'QIN':('Q','Q'),'BOS':('BSP','BSP'),'AVE':('AVP','AVP')}[cid]
 cs.cell(i,16,primary[0]);cs.cell(i,17,urls[primary[1]])
 if cid=='BIO':cs.cell(i,8,'Low (AF recording validation only)')
# Methods reflect new verification while baseline claims remain visibly labeled.
ms=w['Methods']
for row in ms.iter_rows(min_row=2):
 if row[0].value=='Purpose':row[1].value='Original extraction plus targeted primary-study verification on 2026-10-28. Original claims retained; V-prefixed rows add primary evidence, including explicitly tagged later reports.'
 if row[0].value=='AVEIR grading':row[1].value='Primary 2023 study supplies numeric safety 90.3% (271/300); primary safety evidence Low. Original unquantified review claim stays Very low. Primary performance estimates use multiple imputation.'
 if row[0].value=='Verification limits':row[1].value='Nine primary reports checked: six full reports and three original abstracts via author institutions. Two current registry records checked. No IPD audit, exhaustive search, registry history/version audit, or full supplement/SAP audit.'
 if row[0].value=='AI grading':row[1].value='SmartECG includes AF AI and recorded-signal validation; combined signal processing/AI effect, no AI-only clinical benefit trial. Other candidates do not establish ML-specific clinical benefit in checked primary reports.'
ms.append(['Primary-only filter','Filter Evidence layer to Primary-study verification; use Actual model, Timing and Verification status. Earlier generation and model/simulation evidence are not direct candidate clinical trials.'])
ms.append(['Publication timing','February 2024 cover date is reference boundary. BIP published 2024-01-23 (after reference retrieval Jan16); BSP May2024; ALG and LON in 2025 are later evidence, not missing prepublication evidence.'])
ms.append(['Quality versus verification','Confirming numbers/publication does not increase certainty by itself. Added quantitative primary AVEIR safety is Low; Boston primary endpoints Low only in later-evidence layer; model extrapolations Very low for clinical outcomes.'])
ms.append(['Registry caution','Current registry postings differ from publication imputation populations and cutoff dates. Registry checks sheet preserves differences; no unverified inference about original preregistration timing or protocol changes.'])
isu=w['Issues']
isu.append(['I11','AVEIR primary versus registry','AVP and RAV','Different endpoint analysis populations and estimates','Publication imputed 90.2%/97.3%; registry 90.8%/98.2% retained separately.'])
isu.append(['I12','MODULAR primary versus registry','BSP and RBS','Communication test counts differ; enrollment cutoff differs','293 vs297 enrollment due to cutoff; communication 147 vs149 unresolved; same modeled 98.8%.'])
isu.append(['I13','SmartECG scope','BIP methods/results','2024 study uses IIIm proxy and Balanced setting, not 86%/98% release dataset','Do not describe as independent prospective IV replication; industry funded; algorithm on remote server.'])
isu.append(['I14','Micra generation mismatch','MAR; ACC; IDE','Earlier generation clinical studies are not AV2/VR2 trials','Keep actual models and transfer certainty explicit.'])
isu.append(['I15','Post-review evidence','BSP; LON; ALG','2024–2025 results unavailable at February2024 cover date','Tag as later evidence; do not retroactively validate review.'])
# Consistent readable formatting and resizable Excel Tables.
for s in w:
 s.freeze_panes='C2';s.sheet_view.zoomScale=80
 for tab in s.tables.values():
  tab.ref=s.dimensions
  tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
  # openpyxl rebuilds missing column metadata for newly appended columns on save.
  if len(tab.tableColumns)!=s.max_column:tab.tableColumns=[]
 for cell in s[1]:cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D');cell.alignment=Alignment(wrap_text=True,vertical='center')
 s.row_dimensions[1].height=48
 for row in s.iter_rows(min_row=2):
  s.row_dimensions[row[0].row].height=110 if s.title in ['Evidence','Primary trials','Registry checks','Candidates'] else 70
  for c in row:
   c.font=Font(name='Arial',size=10);c.alignment=Alignment(wrap_text=True,vertical='top')
   if isinstance(c.value,str) and c.value.startswith('https://'):c.hyperlink=c.value;c.font=Font(name='Arial',size=10,color='0563C1',underline='single')
   if isinstance(c.value,str) and c.value.startswith(('Very low','Low','Not demonstrated')):c.fill=PatternFill('solid',fgColor='FCE4D6' if c.value.startswith('Very low') else 'FFF2CC')
 for col in range(1,s.max_column+1):s.column_dimensions[get_column_letter(col)].width=20 if 'ID' in str(s.cell(1,col).value) else 38
 s.auto_filter.ref=s.dimensions
for row in ws.iter_rows(min_row=26):
 row[3].comment=Comment(f'Source: {row[16].value}; {row[15].value}; {row[17].value}; checked 2026-10-28','Verification')
w.save(p/'implantable_device_evidence_verified.xlsx')
with (p/'implantable_device_evidence_verified.csv').open('w',newline='',encoding='utf-8-sig') as f:
 writer=csv.writer(f);writer.writerows([[c.value for c in row] for row in ws])
check=load_workbook(p/'implantable_device_evidence_verified.xlsx')
assert len(check['Primary trials']['A'])==10
assert check['Evidence'].max_row==25+len(new)
assert len(set(r[0].value for r in check['Evidence'].iter_rows(min_row=2)))==24+len(new)
assert all(c.data_type!='e' for s in check for row in s for c in row)
assert all(t.ref==s.dimensions for s in check for t in s.tables.values())
print('Verified outputs:',len(trials),'primary reports,',len(new),'new endpoint rows,',24+len(new),'total evidence rows;',len(check.sheetnames),'sheets; no Excel errors.')
