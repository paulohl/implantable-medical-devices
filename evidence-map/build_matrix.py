from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

out=Path('user/workspace/co-scientist-outputs/implantable-device-evidence');out.mkdir(parents=True,exist_ok=True)
P='https://doi.org/10.5121/ijcsit.2024.02101'
B='https://www.prnewswire.com/news-releases/biotronik-announced-first-implant-of-new-implantable-cardiac-monitor-with-artificial-intelligence-targeting-efficiency-in-monitoring-for-cardiac-arrhythmias-301848638.html'
M='https://www.dicardiology.com/content/medtronic-receives-fda-approval-its-next-generation-micra-leadless-pacing-systems'
Q='https://pmc.ncbi.nlm.nih.gov/articles/PMC5126155/'
A='https://abbott.mediaroom.com/2023-05-20-Abbott-Unveils-Late-Breaking-Data-Showing-the-Worlds-First-Dual-Chamber-Leadless-Pacemaker-Study-Met-All-Three-Primary-Endpoints'
S='https://www.implantable-device.com/2021/12/13/boston-scientific-starts-trial-of-emblem-s-icd-with-empower-leadless-pacemaker/'
sources=[['P', 'Leocadio (2024), supplied review', P, 'Full supplied PDF, 18 pages', 'Narrative review; p.15 states no data other than public citations'],['B','BIOTRONIK (2023-06-12), manufacturer release',B,'Full web text checked','AI applies only to AF detection; data on file; Very Specific filter'],['M','Wainscot Media (2023-05-01), manufacturer announcement coverage',M,'Full web text checked','Median projected longevity; not observed 16–17-year follow-up'],['Q','Xiang et al. (2016), randomized noninferiority trial',Q,'Full HTML including methods/results checked','Correct DOI: 10.4103/0366-6999.193443; QM7211 is a ventricular lead'],['A','Abbott (2023-05-20), manufacturer trial-results release',A,'Full web text checked','Prospective single-arm study; 300 patients; 55 centers; three months'],['S','Prutchi (2021-12-13), trial-initiation report',S,'Full web text checked','Review dates it 2023; linked article is dated 2021; no trial results']]
headers=['Candidate ID','Candidate / manufacturer','Function stratum','Population / purpose','Capabilities reported','Metric(s) available','Device-evidence grade','AI-evidence grade','AI interpretation','Conditions / settings','Main limitations / conflicts','Review location','Source IDs','Source URL']
candidates=[
['BIO','BIOMONITOR IV / BIOTRONIK','Diagnostic monitoring','Arrhythmia monitoring; not pacing or defibrillation','SmartECG plus AI filtering; remote monitoring; PAC/PVC distinction','False-positive reduction; true-episode preservation (external cited source)','Very low','Very low','AI explicitly identified in source B, only for AF detection; filtering in Home Monitoring Service Center, not proof of onboard AI','AF filter = Very Specific; comparator = without SmartECG','Data on file, no sample size or uncertainty; combined SmartECG + AI effect cannot isolate AI increment','pp.4,9–10; refs pp.15–16','P; B',B],
['MAV','Micra AV2 / Medtronic','Leadless pacing','Bradycardia; ventricular pacing with atrial mechanical sensing','Leadless; accelerometer atrial sensing; automatic AV synchrony; remote monitoring claimed','Projected longevity; mass; claimed heart-rate operating range','Very low','Not demonstrated','Advanced automatic algorithms described; ML model/training/AI-specific benefit not established','Review: 80–100 beats/min synchrony context; pulse settings and pacing burden NR','Review reports 8–13 years and nearly 16 years; projection not observed survival; source M says tracking upper limit 135 beats/min (different metric)','pp.4–6,9–10','P; M',M],
['MVR','Micra VR2 / Medtronic','Leadless pacing','Bradycardia; ventricular pacing','Leadless pacing; easier programming; remote monitoring claimed','Projected longevity; mass; battery-life improvement claim','Very low','Not demonstrated','AI label in review; no ML model or AI-specific validation in checked sources','Pacing burden, pulse amplitude/width and impedance NR','Review reports 12 years and nearly 17 years; no clinical head-to-head outcome data for these models extracted','pp.4–6,9–10','P; M',M],
['QIN','Qinming8631 DR / Qinming Medical','Transvenous dual-chamber pacing','Bradyarrhythmia; SSS or AV block; trial ages 18–85','Dual-chamber rate response; RF programming described in review','6-month effective pacing versus Talos DR; FAS and PPS','Low','Not demonstrated','Rate response and programmability are not evidence of ML/AI; cited trial does not test AI','Trial: 0.5 ms assessment pulse width; RA threshold ≤1.5 V, RV ≤1.0 V; see evidence rows','QM7211 in review is ventricular lead, not pacemaker; missing endpoints/attrition; source inconsistencies in difference estimates; short follow-up','pp.6–7,9–10; ref p.18','P; Q',Q],
['BOS','EMBLEM S-ICD + EMPOWER / Boston Scientific','Modular defibrillation + leadless pacing','Ventricular tachyarrhythmia protection; bradycardia pacing and ATP','S-ICD plus communicating leadless pacing module','No numerical clinical performance result in review or checked initiation report','Very low','Not demonstrated','Device coordination is described; no AI-specific validation identified','ATP settings, shock energy, detection thresholds and follow-up NR','Review mixes component and system properties: outside-heart claim applies to S-ICD, not leadless pacemaker; cannot defibrillate applies to pacemaker, not combined system','p.7,9–10; ref p.17','P; S',S],
['AVE','AVEIR DR i2i / Abbott','Dual-chamber leadless pacing','Slow/irregular heart rate; dual-chamber pacing','Implant-to-implant communication; synchronous atrial/ventricular pacing','Implant success; successful AV synchrony; activity-specific mean synchrony','Low','Not demonstrated','i2i communication and adaptive pacing described; no ML model or AI-isolating comparator','First 300 patients; 55 centers; three months; seven postures/walking speeds','Single-arm; manufacturer reporting; exact synchrony threshold and endpoint denominators not stated in review/release; no superiority inference','p.8–10; refs pp.15,17','P; A',A]
]
eh=['Evidence ID','Candidate ID','Endpoint / capability','Reported value / wording','Relation','Numeric lower / value','Numeric upper','Unit','Evidence layer','Sample / denominator','Follow-up','Comparator','Settings / conditions','GRADE-lite','Grade reason','Review location / source section','Source ID','Source URL','Interpretation / normalization note']
e=[]
def add(cid,metric,raw,rel='',lo=None,hi=None,unit='',layer='Review extraction',n='NR',follow='NR',comp='None reported',settings='NR',grade='Very low',reason='Uncontrolled secondary/manufacturer claim; methods or uncertainty absent',loc='',sid='P',url=P,note=''):
 e.append([f'E{len(e)+1:02}',cid,metric,raw,rel,lo,hi,unit,layer,n,follow,comp,settings,grade,reason,loc,sid,url,note])
add('BIO','Monitoring capability','Efficient and accurate arrhythmia monitoring',loc='p.4 §3.2; p.9 §5.1',note='No sensitivity, specificity, AUROC or false-alert rate supplied in review')
add('BIO','False-positive detection reduction','86%', '=',86,unit='%',layer='Cited-source supplement',comp='Without SmartECG',settings='AF filter Very Specific in Home Monitoring Service Center; AI only for AF',loc='B: main text and footnotes',sid='B',url=B,note='Relative reduction; not 86% specificity; combined SmartECG+AI claim; not printed numerically in review')
add('BIO','True-episode preservation','98%', '=',98,unit='%',layer='Cited-source supplement',settings='Same filter and AI scope as E02',loc='B: main text and footnotes',sid='B',url=B,note='Preservation claim; do not relabel as clinical sensitivity without validation methods')
add('MAV','Battery longevity projection','nearly 16 years','approximately',16,unit='years',loc='p.5 §4.4.1; p.10 §5.3',note='Confirmed as median projected longevity in M; assumptions NR; not observed durability')
add('MAV','Battery life, conflicting specification','8–13 years','range',8,13,'years',loc='p.5 §4.4.4',note='Retained as conflicting review claim; do not average with E04')
add('MVR','Battery longevity projection','nearly 17 years','approximately',17,unit='years',loc='p.10 §5.3',note='Confirmed as median projected longevity in M; not observed durability')
add('MVR','Battery life, conflicting specification','12 years','=',12,unit='years',loc='p.5 §4.4.4',note='Retained as conflicting review claim; do not average with E06')
for cid in ['MAV','MVR']:
 add(cid,'Battery-life improvement','approximately 40% improvement','approximately',40,unit='%',comp='Predecessor Micra generation, unspecified in review',loc='p.10 §5.3',note='Relative improvement in projected battery life; not patient outcome')
 add(cid,'Device mass','1.75 grams','=',1.75,unit='g',loc='p.10 §5.3',note='Shared statement for AV2 and VR2; not an independent clinical endpoint')
add('MAV','Heart-rate range for claimed improved AV synchrony','80–100 bpm','range',80,100,'beats/min',loc='p.6 §4.4.4',note='Operating context, not measured % synchrony or tracking upper limit; M separately states upper limit 135 beats/min')
add('MAV','Projected one-device-for-life proportion','80% of patients only need one device for life','=',80,unit='%',loc='p.5 §4.4.1',note='Projection, not lifetime follow-up; M states >80% for Micra recipients, not model-specific observed result')
add('QIN','RF-powered programming','Program control without consuming battery power',loc='p.6 §4.5.2',note='Review claim; energy consumption not quantified')
add('QIN','Effective pacing, full analysis set','81.4% vs 79.5%','=',81.4,unit='%',layer='Cited-source supplement',n='113 Qinming; 112 Talos DR (225 total)',follow='6 months',comp='Talos DR: 79.5%',settings='14-center randomized single-blind noninferiority trial; margin −9.5 percentage points; missing values LOCF then non-effective',grade='Low',reason='RCT starts High; −1 missing endpoints/attrition; −1 imprecision/reporting inconsistencies and wide noninferiority margin',loc='Q: Abstract; Methods; Results primary endpoints',sid='Q',url=Q,note='Acceptable capture threshold/sensing composite; not % beats paced. Published difference estimates inconsistent with raw proportions; not propagated')
add('QIN','Effective pacing, per protocol','95.4% vs 89.5%','=',95.4,unit='%',layer='Cited-source supplement',n='87 Qinming; 86 Talos DR',follow='6 months',comp='Talos DR: 89.5%',settings='Same trial; complete-protocol subset',grade='Low',reason='Post-randomization exclusions/missing endpoints; reporting inconsistencies; short outcome horizon',loc='Q: Abstract; Results primary endpoints',sid='Q',url=Q,note='Do not combine PPS and FAS; noninferiority is not superiority')
add('QIN','Capture-threshold assessment pulse width','0.5 ms','=',0.5,unit='ms',layer='Cited-source supplement',settings='Primary-endpoint methods; long-term threshold allows upward fluctuation ≤20%',grade='Not applicable',reason='Protocol setting, not clinical effect estimate',loc='Q: Methods, Efficacy assessments',sid='Q',url=Q,note='Figure caption separately says 0.4 ms; retain methods value, flag inconsistency')
add('QIN','Acceptable atrial capture threshold','≤1.5 V','<=',1.5,unit='V',layer='Cited-source supplement',settings='0.5 ms; output tested at twice measured capture threshold',grade='Not applicable',reason='Protocol acceptance criterion',loc='Q: Methods, Primary end points',sid='Q',url=Q)
add('QIN','Acceptable ventricular capture threshold','≤1.0 V','<=',1,unit='V',layer='Cited-source supplement',settings='0.5 ms; output tested at twice measured capture threshold',grade='Not applicable',reason='Protocol acceptance criterion',loc='Q: Methods, Primary end points',sid='Q',url=Q)
add('BOS','Coordinated pacing/defibrillation capability','S-ICD can trigger leadless pacemaker ATP',loc='p.7 §4.6.1–4.6.4',note='No implant success, ATP efficacy, shock success, adverse-event rate or numeric settings extracted; S confirms trial initiation only')
for metric,raw,rel,val,note in [
 ('Implant success','98.3%','=',98.3,'Procedural success; denominator-specific event count not reconstructed from rounded percentage'),
 ('Patients with successful AV synchrony','more than 97%','>',97,'Proportion of patients meeting a synchrony criterion, not average percentage of synchronous beats; exact criterion NR in review/release'),
 ('Mean AV synchrony in each activity condition','more than 95%','>',95,'Mean synchrony in each of seven postures/walking speeds; not interchangeable with proportion of successful patients')]:
 add('AVE',metric,raw,rel,val,unit='%',n='First 300 enrolled; endpoint-specific denominator NR',follow='3 months',settings='Prospective international single-arm; 55 centers; activity analysis includes seven postures/walking speeds',grade='Low',reason='Single-arm prospective study starts Low; short horizon, manufacturer summary; no comparative/AI inference',loc='p.8 §4.7.2',sid='P; A',url=A,note=note)
add('AVE','Safety endpoint','Freedom from device or procedure complications endpoint achieved',n='First 300 enrolled; endpoint-specific denominator NR',follow='3 months',grade='Very low',reason='Pass/fail manufacturer summary without numeric event rate or uncertainty in review/release',loc='p.8 §4.7.2',sid='P; A',url=A,note='Not 100% complication-free; exact safety rate not supplied in these extracted sources')

methods=[
['Purpose','Paper-bounded evidence map with targeted checks of cited sources, not a systematic review or treatment ranking. Extracted 2026-10-28; supplements published before the 2024 review.'],
['Scope','Six candidate rows across five manufacturers: split Micra AV2 and VR2; retain EMBLEM+EMPOWER as a combined system. Talos DR appears only as the Qinming trial comparator.'],
['Read first','Candidate matrix is the overview. Evidence rows are atomic claims; filter by endpoint, evidence layer and grade before sorting numeric values.'],
['Adapted recipe','No candidate drugs or dosing regimens. Dose is not applicable. Device settings are extracted only where reported; pulse amplitudes, shock energy and algorithm thresholds are NR elsewhere.'],
['Evidence layers','Review extraction = claim printed in supplied review. Cited-source supplement = additional numeric detail from its references; never silently attributed to the review.'],
['Units','Percentages remain numeric values on a 0–100 scale (98.3 means 98.3%); years, months, g, beats/min, V and ms. No unsupported conversion of 3 months to 90 days. No derived effect sizes or pooled score.'],
['Relations','Keep >, <=, approximately and range separately. Bounds are not point estimates. Blank numeric cell = nonnumeric/not reported, never zero.'],
['GRADE-lite','Endpoint-specific pragmatic certainty, not formal GRADE or regulatory assessment. High=4, Moderate=3, Low=2, Very low=1; Not demonstrated is absence of extracted AI validation, not evidence of no AI.'],
['Starting certainty','Randomized direct comparative clinical evidence starts High; prospective single-arm evidence starts Low; manufacturer/data-on-file/narrative capability claims start Very low.'],
['Downgrades','One level for each serious concern: bias/missing endpoints, inconsistency, indirectness, imprecision, selective reporting/publication. Floor Very low; no upgrades. Protocol specifications are Not applicable.'],
['Qinming grading','High minus one for missing endpoints/attrition and minus one for imprecision/reporting limitations = Low for six-month effective pacing. The RCT does not establish long-term safety or AI benefit.'],
['AVEIR grading','Low for short-term descriptive implant/synchrony endpoints; no head-to-head comparator. Safety pass/fail claim Very low because the extracted review/release does not supply an event rate.'],
['AI grading','No extracted study isolates an AI-versus-no-AI clinical benefit. BIOMONITOR explicitly uses AI but its numeric claim is combined SmartECG+AI, data on file. Other candidates: Not demonstrated in checked materials.'],
['Candidate grades','Overview grade summarizes the most informative extracted device-performance evidence, not every capability. Use evidence-row grade for the actual claim.'],
['No ranking','Monitoring, bradycardia pacing and defibrillation differ in indication and outcome. Projected years, diagnostic filtering and clinical endpoint percentages must not share a performance leaderboard.'],
['Missing and conflict flags','NR=not reported in extracted source. Conflicting longevity/specifications remain separate rows. AI model architecture, training dataset, external validation and inference location mostly unreported.'],
['Safety','Review lists possible complications without rates for Micra and Boston. No cross-device safety ranking; no absence-of-harm inference from unreported data. Not a clinical recommendation.'],
['Excluded contextual material','Electronic nose network, neural-implant review, microneedle diabetes paper, germline ethics and generic IoT sources lack candidate-specific implant comparison results in this review. BACKBEAT citation is mismatched to AVEIR discussion, not an extra evaluated candidate.'],
['Verification limits','Five linked cited sources checked; no exhaustive primary-study search, current approval-status audit, independent data-on-file audit, or systematic risk-of-bias instrument. AVEIR NEJM search result not used for extra numeric endpoints.']]
issues=[['I01','Micra AV2','pp.5,10','8–13 years versus nearly 16 years','Retain both; M identifies nearly 16 as median projected longevity.'],['I02','Micra VR2','pp.5,10','12 years versus nearly 17 years','Retain both; M identifies nearly 17 as median projected longevity.'],['I03','Qinming model','pp.6,9–10','8631 DR versus QM7211','Q methods identify QM7211 as ventricular lead. Use Qinming8631 DR as candidate.'],['I04','Boston system','p.7; p.10','Entire system said to stay outside heart; cannot defibrillate','Separate S-ICD properties from intracardiac leadless pacemaker; combined system includes defibrillation.'],['I05','Boston reference','p.17','2023 citation links to 2021 trial initiation','Use checked article date 2021-12-13; no results inferred.'],['I06','AVEIR','pp.8,17','BACKBEAT reference attached to AVEIR description','Use Abbott release for AVEIR study; do not import BACKBEAT evidence.'],['I07','AI label','Device sections','Automatic algorithms/communication called AI without model validation','Keep device evidence distinct from AI-specific evidence.'],['I08','Qinming primary study','Q Results and Methods','Reported differences do not consistently match raw proportions; 0.5 ms methods vs 0.4 ms figure caption','Retain group endpoint percentages; do not calculate superiority or propagate inconsistent effect estimates.'],['I09','AVEIR endpoints','p.8','>97% patients successful vs >95% mean synchrony','Separate denominators and estimands; neither is a head-to-head advantage.'],['I10','BIOMONITOR numeric claims','B footnote','86%/98% depend on Very Specific filter; AI AF-only','Not generic specificity/sensitivity; comparator lacks SmartECG, so AI-only increment unresolved.']]

wb=Workbook();wb.remove(wb.active)
def sheet(name,heads,rows):
 ws=wb.create_sheet(name);ws.append(heads)
 for row in rows:ws.append(row)
 ws.freeze_panes='C2';ws.sheet_view.zoomScale=85
 tab=Table(displayName=name.replace(' ','')+'Table',ref=ws.dimensions);tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(tab)
 for c in ws[1]: c.font=Font(name='Arial',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D');c.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=42
 for row in ws.iter_rows(min_row=2):
  ws.row_dimensions[row[0].row].height=100 if name in ['Candidates','Evidence'] else 66
  for c in row:
   c.font=Font(name='Arial',size=10);c.alignment=Alignment(wrap_text=True,vertical='top')
   if isinstance(c.value,str) and c.value.startswith('https://'):c.hyperlink=c.value;c.font=Font(name='Arial',size=10,color='0563C1',underline='single')
   if c.value in ['Low','Very low','Not demonstrated']:
    c.fill=PatternFill('solid',fgColor={'Low':'FFF2CC','Very low':'FCE4D6','Not demonstrated':'E7E6E6'}[c.value])
 for i,h in enumerate(heads,1):ws.column_dimensions[get_column_letter(i)].width=20 if ('ID' in h or 'grade' in h.lower()) else 38
 ws.auto_filter.ref=ws.dimensions
 return ws
sheet('Candidates',headers,candidates)
ws=sheet('Evidence',eh,e)
for row in range(2,ws.max_row+1):
 ws.cell(row,6).number_format='0.###';ws.cell(row,7).number_format='0.###'
 ws.cell(row,4).comment=Comment('Source: '+str(ws.cell(row,17).value)+'; '+str(ws.cell(row,16).value)+'; '+str(ws.cell(row,18).value),'Source')
sheet('Methods',['Item','Definition / rule'],methods)
sheet('Issues',['Issue ID','Subject','Location','Issue','Treatment'],issues)
sheet('Sources',['Source ID','Source','URL','Access / verification','Notes'],sources)
wb.save(out/'implantable_device_evidence.xlsx')
with (out/'implantable_device_evidence.csv').open('w',newline='',encoding='utf-8-sig') as f:
 w=csv.writer(f);w.writerow(eh);w.writerows(e)
report='''# Implantable-device evidence map

**Bottom line:** the paper supports a function-stratified evidence map, not a ranking of “best AI implants.” It is a narrative review with no original dataset (p.15). Six candidate systems/models across five manufacturers were extracted. No reviewed comparison isolates the clinical contribution of AI.

## What the comparison shows

| Candidate | Most useful extracted evidence | Certainty and boundary |
|---|---|---|
| BIOMONITOR IV | Cited manufacturer release: 86% fewer false-positive detections while preserving 98% of true episodes | Very low; data on file, AF filter “Very Specific”; AI applies only to AF. Comparator is without SmartECG, not an AI-only ablation. |
| Micra AV2 | Nearly 16-year median projected longevity; review also says 8–13 years | Very low for longevity claim; projection, not observed long-term follow-up. |
| Micra VR2 | Nearly 17-year median projected longevity; review also says 12 years | Very low for longevity claim; conflicting review specifications retained. |
| Qinming8631 DR | Cited randomized trial: effective pacing 81.4% versus 79.5% with Talos DR in full analysis set, six months | Low for this endpoint; missing assessments/attrition and reporting limitations. Noninferiority does not establish superiority or an AI benefit. |
| EMBLEM S-ICD + EMPOWER | Coordinated defibrillation and leadless pacing/ATP described | Very low; no numerical clinical performance result in review or checked trial-initiation report. |
| AVEIR DR | 98.3% implant success; >97% of patients with successful AV synchrony; mean synchrony >95% in each of seven activity conditions | Low for short-term descriptive performance; single arm, 300 initially enrolled, 55 centers, three months. These percentages measure different things. |

## Corrections that matter

- **Qinming identity:** QM7211 is the ventricular lead in the cited trial, not the dual-chamber pacemaker. The candidate is Qinming8631 DR.
- **Boston component boundaries:** the S-ICD is extrathoracic; the leadless pacemaker is intracardiac. “Cannot defibrillate” applies to the pacemaker component, not the combined system.
- **Longevity:** the Micra figures are projections. The review’s conflicting values remain visible rather than being averaged or silently replaced.
- **AI versus automation:** automatic synchrony, rate response and implant-to-implant communication do not by themselves establish a machine-learning implementation or its added clinical benefit.

## How to use the files

The Excel workbook has five filterable sheets: **Candidates**, **Evidence**, **Methods**, **Issues**, and **Sources**. Start with Candidates; use Evidence for atomic claims and exact source conditions. The CSV is the same long-form Evidence sheet for reuse. Source IDs and URLs distinguish claims printed in the supplied review from extra details recovered from its cited sources. Page references are printed PDF page numbers.

Percentages use a 0–100 numeric scale. Approximate values and inequalities retain their qualifiers. Units are years, months, g, beats/min, V and ms. No drugs or dose regimens are compared, so the recipe’s dose normalization was replaced with device-setting extraction. Missing settings are **NR**, not zero; protocol acceptance thresholds are not performance outcomes.

## GRADE-lite and limits

This is a pragmatic endpoint-level certainty assessment, not formal GRADE. Randomized evidence starts High, prospective single-arm evidence Low, and narrative/manufacturer/data-on-file claims Very low. Serious bias, inconsistency, indirectness, imprecision or selective reporting can lower certainty, with a Very low floor. Qinming’s six-month pacing endpoint is downgraded twice from High to Low; its trial does not test AI. AVEIR’s short-term single-arm performance remains Low, while its unquantified safety pass/fail claim is Very low. Protocol settings are Not applicable. “Not demonstrated” means AI validation was not established in the checked materials, not that the device has no AI.

No numeric overall score, cross-indication safety ranking, pooled effect estimate, or recommendation is justified. Five cited web sources were checked, including the Qinming full trial; this is not an exhaustive search or an audit of current regulatory status. The Qinming source itself contains inconsistencies in reported between-group differences and pulse-width descriptions; the workbook flags them and does not propagate questionable effect estimates. No synthetic patient data were used.

## Checked sources

'''
for sid,title,url,access,note in sources:report+=f'- **{sid}: {title}** — {url}\n'
(out/'evidence_readme.md').write_text(report,encoding='utf-8')
check=load_workbook(out/'implantable_device_evidence.xlsx')
assert check['Candidates'].max_row==7
assert check['Evidence'].max_row==len(e)+1
assert all(r[1] in {c[0] for c in candidates} for r in e)
assert len({r[0] for r in e})==len(e)
assert all(len(r)==len(eh) for r in e)
assert all(c.data_type!='e' for ws in check for row in ws for c in row)
print(f'Validated: {len(candidates)} candidates, {len(e)} evidence rows, {len(issues)} issues, {len(sources)} sources; five Excel tables; no formulas or error cells.')
